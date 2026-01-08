# import PyPDF2
# import docx
# import pandas as pd
# import json
# import os

# def process_file(file_path: str) -> dict:
#     if file_path.endswith(".pdf"):
#         with open(file_path, "rb") as f:
#             reader = PyPDF2.PdfReader(f)
#             text = "".join(page.extract_text() for page in reader.pages)
#         return {"text": text}
#     elif file_path.endswith(".docx"):
#         doc = docx.Document(file_path)
#         text = "\n".join(para.text for para in doc.paragraphs)
#         return {"text": text}
#     elif file_path.endswith(".csv"):
#         df = pd.read_csv(file_path)
#         return {"data": df.to_dict()}
#     elif file_path.endswith(".json"):
#         with open(file_path, "r") as f:
#             data = json.load(f)
#         return {"data": data}
#     elif file_path.endswith(".txt"):
#         with open(file_path, "r") as f:
#             text = f.read()
#         return {"text": text}
#     return {"error": "Unsupported file type"}


import PyPDF2
import docx
import pandas as pd
import json
import os
import re
from langchain_text_splitters import RecursiveCharacterTextSplitter

def process_file(file_path: str) -> list[str]:
    """Extract text and split into smaller, smarter chunks."""
    if not os.path.exists(file_path):
        raise ValueError(f"File not found: {file_path}")
    
    file_ext = os.path.splitext(file_path)[1].lower()
    text = ""
    
    if file_ext == ".pdf":
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
    
    elif file_ext == ".docx":
        doc = docx.Document(file_path)
        text = "\n".join(para.text for para in doc.paragraphs)
    
    elif file_ext in [".csv", ".json"]:
        if file_ext == ".csv":
            df = pd.read_csv(file_path)
            text = df.to_string(index=False)
        else:
            with open(file_path, "r") as f:
                data = json.load(f)
            text = json.dumps(data, indent=2)
    
    elif file_ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
    
    else:
        raise ValueError(f"Unsupported file type: {file_ext}")
    
    if not text.strip():
        raise ValueError("No text extracted from file.")
    
    # === NEW: Smaller, smarter chunks ===
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=300,          # Smaller chunks → better precision
        chunk_overlap=50,        # Overlap → preserves context across sections
        separators=["\n\n", "\n", ". ", "? ", "! ", " "],  # Respects paragraphs/sentences
        length_function=len,
    )
    chunks = splitter.split_text(text)
    
    # Filter tiny chunks
    chunks = [chunk.strip() for chunk in chunks if len(chunk.strip()) > 30]
    
    if not chunks:
        raise ValueError("No valid chunks generated.")
    
    return chunks